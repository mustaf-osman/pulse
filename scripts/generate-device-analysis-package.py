from __future__ import annotations

import csv
import html
import os
import textwrap
import zipfile
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "device-integration-analysis"
GENERATED_AT = datetime.now().strftime("%Y-%m-%d %H:%M")

SUMMARY = [
    "Nimo 设备中心 v1 已完成软件侧硬件接入准备：设备 API、设备中心 UI、模拟设备闭环、局域网启动提示、固件请求示例和 smoke 自检。",
    "当前软件处于可对接真实 ESP32-S3 固件的阶段；下一步重点从功能可用转向安全、持久化、诊断和长期维护。",
    "短期硬件接入建议采用 HTTP 轮询方案：开机注册、定时心跳、读取配置、轮询下一条提醒、上报完成/稍后。",
]

CURRENT_STATUS = [
    ["模块", "状态", "说明"],
    ["设备中心 UI", "已完成", "顶部/左侧导航已加入“设备”，可展示接入地址、LAN 状态、设备列表、提醒、协议和固件示例。"],
    ["设备 API", "已完成", "已具备 register、heartbeat、config、status、next-reminder、complete、snooze、test-push、stream。"],
    ["模拟设备", "已完成", "可在无硬件条件下验证注册、心跳、测试推送、完成、稍后。"],
    ["协议自检", "已完成", "新增 npm run smoke:device，非破坏性检查设备协议。"],
    ["LAN 模式", "已具备", "已有 npm run start:lan；硬件接入时需电脑和 ESP32-S3 处于同一 Wi-Fi。"],
    ["安全配对", "待增强", "建议新增 pairing code、deviceToken 和接口鉴权。"],
    ["设备落库", "待增强", "当前设备注册更偏运行态，建议落库保存真实设备、token、固件版本、最后心跳。"],
    ["设备事件日志", "待增强", "建议记录硬件请求和错误，方便调试 ESP32-S3。"],
]

PROTOCOL = [
    ["接口", "方法", "用途", "请求参数/Body", "响应重点", "是否会改数据"],
    ["/device/info", "GET", "获取服务器、设备列表和协议清单", "无", "server、devices、protocol", "否"],
    ["/device/register", "POST", "设备开机注册", "id、name、type、capabilities", "device、server", "是，登记/更新设备"],
    ["/device/heartbeat", "POST", "设备心跳保活", "deviceId", "device、serverTime", "是，更新 lastSeenAt"],
    ["/device/config", "GET", "读取设备配置", "deviceId 查询参数", "pollIntervalSeconds、heartbeatSeconds、actions、snoozeMinutes", "轻微，touch 设备"],
    ["/device/status", "GET", "读取设备状态", "deviceId 可选", "server、devices", "轻微，touch 设备"],
    ["/device/next-reminder", "GET", "拉取下一条 active 提醒", "deviceId 查询参数", "reminder 或 null", "轻微，touch 设备"],
    ["/device/complete", "POST", "硬件完成提醒", "deviceId、reminderId", "changes、reminder", "是，完成提醒"],
    ["/device/snooze", "POST", "硬件稍后提醒", "deviceId、reminderId、minutes", "dueAt、minutes、reminder", "是，修改提醒时间"],
    ["/device/test-push", "POST", "测试推送事件", "deviceId、title、task 可选", "payload", "否，不改真实提醒"],
    ["/device/stream", "GET", "设备 SSE 实时事件流", "deviceId 查询参数", "SSE 事件", "否"],
]

FIRMWARE_FLOW = [
    ["步骤", "固件动作", "软件接口", "成功标准"],
    ["1", "连接 Wi-Fi", "无", "ESP32-S3 与电脑在同一局域网"],
    ["2", "开机注册设备", "POST /device/register", "设备中心显示新设备"],
    ["3", "读取配置", "GET /device/config", "获得轮询和心跳间隔"],
    ["4", "定时心跳", "POST /device/heartbeat", "设备中心 90 秒内显示在线"],
    ["5", "轮询提醒", "GET /device/next-reminder", "屏幕显示下一条 active 提醒"],
    ["6", "用户点完成", "POST /device/complete", "提醒状态变为 completed"],
    ["7", "用户点稍后", "POST /device/snooze", "提醒 dueAt 延后指定分钟"],
]

ROADMAP = [
    ["阶段", "目标", "软件任务", "优先级", "验收方式"],
    ["P0 已完成", "硬件接入协议可用", "设备中心、API、模拟设备、smoke:device", "已完成", "npm run smoke:device 通过"],
    ["P1", "真实硬件安全接入", "配对码、deviceToken、接口鉴权", "高", "未授权请求被拒绝，已配对设备可用"],
    ["P1", "设备可长期管理", "设备表落库、重命名、删除、固件版本、最后心跳", "高", "重启后设备仍存在"],
    ["P1", "调试可观测", "设备事件日志、错误统计、最近请求、离线原因", "中高", "设备页可看到最近 20 条事件"],
    ["P2", "提醒展示更精细", "ack、displayText、speakText、priority、sound、requiresAck", "中", "硬件不重复展示已确认提醒"],
    ["P2", "LAN 诊断", "端口检测、防火墙提示、二维码配置", "中", "页面可判断硬件是否可连"],
    ["P3", "长期维护", "OTA 元信息、配置下发、WebSocket/MQTT 方案", "中低", "设备可上报版本并读取升级信息"],
]

RISKS = [
    ["风险", "影响", "当前状态", "建议处理"],
    ["局域网未开启", "ESP32-S3 无法访问电脑 API", "当前默认 127.0.0.1", "硬件调试使用 npm run start:lan，并放行 Windows 防火墙"],
    ["无设备鉴权", "同网段设备可能误调用提醒接口", "待增强", "实现 pairing code + deviceToken"],
    ["设备状态未持久化", "重启后设备记录丢失", "待增强", "新增 devices 表和 device_events 表"],
    ["轮询延迟", "提醒非实时", "可接受", "第一版使用 15 秒轮询；后续评估 SSE/WebSocket/MQTT"],
    ["普通 Node 与 better-sqlite3 ABI 不一致", "部分本地脚本不能直接读提醒库", "已在 smoke 中跳过 DB 依赖项", "使用 Electron runtime 或重建 better-sqlite3"],
    ["防火墙阻断", "硬件无法连接", "需现场验证", "设备中心加入 LAN 诊断和二维码"],
]

FILES_CHANGED = [
    ["文件", "用途"],
    ["src/api.js", "设备 API、服务器信息、LAN 地址、协议清单"],
    ["src/ui/brain-ui/device-view.js", "设备中心前端逻辑、模拟设备、固件示例"],
    ["src/ui/brain-ui/app-shell.js", "设备页面结构、导航入口、接入清单"],
    ["src/ui/brain-ui/app.js", "初始化设备中心、SSE 事件分发"],
    ["src/ui/brain-ui/styles.css", "设备中心样式"],
    ["scripts/smoke-device.mjs", "设备协议非破坏性 smoke test"],
    ["package.json", "新增 npm run smoke:device"],
]


def ensure_out():
    OUT.mkdir(parents=True, exist_ok=True)


def md_table(rows):
    head = rows[0]
    body = rows[1:]
    lines = ["| " + " | ".join(head) + " |", "| " + " | ".join(["---"] * len(head)) + " |"]
    for row in body:
        lines.append("| " + " | ".join(str(cell).replace("|", "\\|") for cell in row) + " |")
    return "\n".join(lines)


def build_markdown():
    sections = [
        "# Nimo 硬件接入软件层分析报告",
        f"生成时间：{GENERATED_AT}",
        "## 一、结论摘要",
        "\n".join(f"- {item}" for item in SUMMARY),
        "## 二、当前完成状态",
        md_table(CURRENT_STATUS),
        "## 三、设备协议 v1",
        md_table(PROTOCOL),
        "## 四、ESP32-S3 固件最小接入流程",
        md_table(FIRMWARE_FLOW),
        "## 五、软件后续路线图",
        md_table(ROADMAP),
        "## 六、风险与建议",
        md_table(RISKS),
        "## 七、涉及文件",
        md_table(FILES_CHANGED),
        "## 八、建议下一步",
        "1. 优先实现设备配对码与 deviceToken。\n2. 新增设备表和设备事件日志表。\n3. 给设备中心增加局域网诊断和二维码配置。\n4. 硬件到手后先用 HTTP 轮询跑通注册、心跳、拉取提醒、完成、稍后。",
    ]
    return "\n\n".join(sections) + "\n"


def write_markdown(content):
    path = OUT / "Nimo_Device_Integration_Analysis.md"
    path.write_text(content, encoding="utf-8")
    return path


def write_html(content):
    body = []
    for block in content.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("# "):
            body.append(f"<h1>{html.escape(block[2:])}</h1>")
        elif block.startswith("## "):
            body.append(f"<h2>{html.escape(block[3:])}</h2>")
        elif block.startswith("| "):
            rows = [line for line in block.splitlines() if line.startswith("| ")]
            parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in rows]
            parsed = [parsed[0]] + parsed[2:]
            table = ["<table>"]
            table.append("<thead><tr>" + "".join(f"<th>{html.escape(cell)}</th>" for cell in parsed[0]) + "</tr></thead>")
            table.append("<tbody>")
            for row in parsed[1:]:
                table.append("<tr>" + "".join(f"<td>{html.escape(cell)}</td>" for cell in row) + "</tr>")
            table.append("</tbody></table>")
            body.append("\n".join(table))
        elif block.startswith("- "):
            items = [line[2:] for line in block.splitlines() if line.startswith("- ")]
            body.append("<ul>" + "".join(f"<li>{html.escape(item)}</li>" for item in items) + "</ul>")
        else:
            body.append("<p>" + html.escape(block).replace("\n", "<br>") + "</p>")
    page = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>Nimo 硬件接入软件层分析报告</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, 'Microsoft YaHei', 'Segoe UI', sans-serif; margin: 42px; color: #111827; line-height: 1.65; }}
h1 {{ font-size: 30px; margin-bottom: 8px; }}
h2 {{ margin-top: 32px; padding-bottom: 6px; border-bottom: 1px solid #e5e7eb; color: #1d4ed8; }}
table {{ border-collapse: collapse; width: 100%; margin: 14px 0 22px; font-size: 13px; }}
th, td {{ border: 1px solid #d1d5db; padding: 8px 10px; vertical-align: top; }}
th {{ background: #eff6ff; color: #1e3a8a; text-align: left; }}
tr:nth-child(even) td {{ background: #f9fafb; }}
code {{ background: #eef2ff; color: #3730a3; padding: 2px 5px; border-radius: 5px; }}
@media print {{ body {{ margin: 18mm; }} h2 {{ page-break-after: avoid; }} table {{ page-break-inside: avoid; }} }}
</style>
</head>
<body>
{''.join(body)}
</body>
</html>"""
    path = OUT / "Nimo_Device_Integration_Analysis.html"
    path.write_text(page, encoding="utf-8")
    return path


def add_docx_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for i, cell in enumerate(rows[0]):
        table.rows[0].cells[i].text = str(cell)
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, cell in enumerate(row):
            cells[i].text = str(cell)


def write_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    doc.add_heading("Nimo 硬件接入软件层分析报告", 0)
    doc.add_paragraph(f"生成时间：{GENERATED_AT}")
    doc.add_heading("一、结论摘要", level=1)
    for item in SUMMARY:
        doc.add_paragraph(item, style="List Bullet")
    for title, rows in [
        ("二、当前完成状态", CURRENT_STATUS),
        ("三、设备协议 v1", PROTOCOL),
        ("四、ESP32-S3 固件最小接入流程", FIRMWARE_FLOW),
        ("五、软件后续路线图", ROADMAP),
        ("六、风险与建议", RISKS),
        ("七、涉及文件", FILES_CHANGED),
    ]:
        doc.add_heading(title, level=1)
        add_docx_table(doc, rows)
    doc.add_heading("八、建议下一步", level=1)
    for item in ["优先实现设备配对码与 deviceToken。", "新增设备表和设备事件日志表。", "给设备中心增加局域网诊断和二维码配置。", "硬件到手后先用 HTTP 轮询跑通注册、心跳、拉取提醒、完成、稍后。"]:
        doc.add_paragraph(item, style="List Number")
    path = OUT / "Nimo_Device_Integration_Analysis.docx"
    doc.save(path)
    return path


def register_pdf_font():
    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simsun.ttc",
    ]
    for font in candidates:
        if font.exists():
            try:
                pdfmetrics.registerFont(TTFont("CNFont", str(font)))
                return "CNFont"
            except Exception:
                pass
    return "Helvetica"


def pdf_paragraph(text, style):
    return Paragraph(html.escape(str(text)).replace("\n", "<br/>"), style)


def write_pdf():
    font = register_pdf_font()
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("CNNormal", parent=styles["Normal"], fontName=font, fontSize=8.5, leading=12, alignment=TA_LEFT)
    title = ParagraphStyle("CNTitle", parent=styles["Title"], fontName=font, fontSize=18, leading=24)
    heading = ParagraphStyle("CNHeading", parent=styles["Heading2"], fontName=font, fontSize=13, leading=18, textColor=colors.HexColor("#1d4ed8"))
    path = OUT / "Nimo_Device_Integration_Analysis.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    story = [pdf_paragraph("Nimo 硬件接入软件层分析报告", title), pdf_paragraph(f"生成时间：{GENERATED_AT}", normal), Spacer(1, 6)]
    story.append(pdf_paragraph("一、结论摘要", heading))
    for item in SUMMARY:
        story.append(pdf_paragraph("• " + item, normal))
    for title_text, rows in [
        ("二、当前完成状态", CURRENT_STATUS),
        ("三、设备协议 v1", PROTOCOL),
        ("四、ESP32-S3 固件最小接入流程", FIRMWARE_FLOW),
        ("五、软件后续路线图", ROADMAP),
        ("六、风险与建议", RISKS),
        ("七、涉及文件", FILES_CHANGED),
    ]:
        story.append(PageBreak())
        story.append(pdf_paragraph(title_text, heading))
        data = [[pdf_paragraph(cell, normal) for cell in row] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff6ff")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(table)
    doc.build(story)
    return path


def write_csv(name, rows):
    path = OUT / name
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)
    return path


def column_name(index):
    result = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def sheet_xml(rows):
    lines = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>']
    for r_idx, row in enumerate(rows, 1):
        lines.append(f'<row r="{r_idx}">')
        for c_idx, value in enumerate(row):
            ref = f"{column_name(c_idx)}{r_idx}"
            value = xml_escape(str(value))
            lines.append(f'<c r="{ref}" t="inlineStr"><is><t>{value}</t></is></c>')
        lines.append('</row>')
    lines.append('</sheetData></worksheet>')
    return "".join(lines)


def write_xlsx():
    sheets = [
        ("当前状态", CURRENT_STATUS),
        ("设备协议", PROTOCOL),
        ("固件流程", FIRMWARE_FLOW),
        ("路线图", ROADMAP),
        ("风险", RISKS),
        ("涉及文件", FILES_CHANGED),
    ]
    path = OUT / "Nimo_Device_Integration_Tables.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>''' + "".join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1, len(sheets) + 1)) + "</Types>")
        z.writestr("_rels/.rels", '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>''')
        z.writestr("xl/_rels/workbook.xml.rels", '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">''' + "".join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(sheets) + 1)) + "</Relationships>")
        sheet_defs = "".join(f'<sheet name="{xml_escape(name)}" sheetId="{i}" r:id="rId{i}"/>' for i, (name, _) in enumerate(sheets, 1))
        z.writestr("xl/workbook.xml", f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>{sheet_defs}</sheets></workbook>''')
        for i, (_, rows) in enumerate(sheets, 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", sheet_xml(rows))
    return path


def main():
    ensure_out()
    md = build_markdown()
    files = [
        write_markdown(md),
        write_html(md),
        write_docx(),
        write_pdf(),
        write_xlsx(),
        write_csv("device_protocol.csv", PROTOCOL),
        write_csv("implementation_roadmap.csv", ROADMAP),
        write_csv("risk_register.csv", RISKS),
    ]
    manifest = OUT / "manifest.txt"
    manifest.write_text("\n".join(str(path.name) for path in files) + "\n", encoding="utf-8")
    files.append(manifest)
    for path in files:
        print(path)


if __name__ == "__main__":
    main()
